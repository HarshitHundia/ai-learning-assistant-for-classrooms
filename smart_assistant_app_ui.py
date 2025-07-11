import streamlit as st
import requests
import json
import os
from transformers import pipeline # For the local Flan-T5 model and emotion classifier
import string # Import string module for punctuation
import pypdf # For reading PDF files
from docx import Document # For reading .docx files
# No need for sounddevice, whisper, pyttsx3, numpy, scipy.io.wavfile, threading, time, msvcrt for Streamlit UI

# --- Configuration & Global Variables ---
# IMPORTANT: For local execution, you MUST provide your own Gemini API key.
# It's recommended to set it as an environment variable (e.g., export GEMINI_API_KEY="your_key")
# or hardcode it for quick testing.
# For local running, we assume the user will set their own key.
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "AIzaSyAgrpayn0eNJy-hqm8OrOOGf6mFrzyaGPE") # Replaced with your actual key

if GEMINI_API_KEY == "YOUR_GEMINI_API_KEY_HERE": # Check if the placeholder is still there
    st.warning("❗ WARNING: Gemini API Key is still the placeholder. Replace it with your actual key for full functionality.")

# --- Real-time Search API Configuration ---
# You MUST replace these with your actual Google Custom Search API Key and Search Engine ID (CX)
# Get them from:
# 1. Google Cloud Console for API Key: https://console.cloud.google.com/apis/credentials
# 2. Programmable Search Engine for CX: https://programmablesearchengine.google.com/controlpanel/all
GOOGLE_CSE_API_KEY = os.getenv("GOOGLE_CSE_API_KEY", "AIzaSyAgrpayn0eNJy-hqm8OrOOGf6mFrzyaGPE") # Replaced with your actual key
GOOGLE_CSE_ID = os.getenv("GOOGLE_CSE_ID", "e16d416d56a8a40b2") # Replaced with your actual CX ID

if GOOGLE_CSE_API_KEY == "YOUR_GOOGLE_CSE_API_KEY_HERE" or GOOGLE_CSE_ID == "YOUR_GOOGLE_CSE_ID_HERE":
    st.warning("❗ WARNING: Google Custom Search API keys are placeholders. Replace them for real search functionality.")

# --- Initialize Session State ---
# Streamlit reruns the script from top to bottom on every interaction.
# st.session_state is used to persist variables across these reruns.
if 'chat_history' not in st.session_state:
    st.session_state.chat_history = []
if 'science_notes_content' not in st.session_state:
    st.session_state.science_notes_content = None
if 'loaded_file_name' not in st.session_state:
    st.session_state.loaded_file_name = None
if 'science_notes_chunks' not in st.session_state:
    st.session_state.science_notes_chunks = []
if 'last_generated_quiz' not in st.session_state:
    st.session_state.last_generated_quiz = None
if 'last_generated_summary' not in st.session_state:
    st.session_state.last_generated_summary = None

# --- Local model setup for contextual QA ---
@st.cache_resource # Cache the model loading to avoid reloading on every rerun
def load_local_qa_generator():
    st.info("⏳ Loading local language model (google/flan-t5-base)... This may take a moment the first time.")
    try:
        qa_generator = pipeline("text2text-generation", model="google/flan-t5-base")
        st.success("✅ Local language model loaded successfully.")
        return qa_generator
    except Exception as e:
        st.error(f"❌ Error loading local model: {e}")
        st.error("Please ensure you have 'transformers' installed (`pip install transformers`) and an active internet connection for the initial download.")
        return None

local_qa_generator = load_local_qa_generator()

# --- Emotion Classifier Setup ---
@st.cache_resource # Cache the model loading
def load_emotion_classifier():
    st.info("⏳ Loading emotion classifier model (j-hartmann/emotion-english-distilroberta-base)...")
    try:
        classifier = pipeline("text-classification", model="j-hartmann/emotion-english-distilroberta-base")
        st.success("✅ Emotion classifier loaded successfully.")
        return classifier
    except Exception as e:
        st.error(f"❌ Error loading emotion classifier: {e}")
        st.error("Please ensure you have 'transformers' installed and an active internet connection for the initial download.")
        return None

emotion_classifier = load_emotion_classifier()

# --- External Tool: Google Search ---
def search_tool(query: str) -> str:
    """
    Performs a real-time web search for the given query using Google Custom Search JSON API.
    Requires GOOGLE_CSE_API_KEY and GOOGLE_CSE_ID to be configured.
    """
    if GOOGLE_CSE_API_KEY == "YOUR_GOOGLE_CSE_API_KEY_HERE" or GOOGLE_CSE_ID == "YOUR_GOOGLE_CSE_ID_HERE":
        st.warning("❗ Warning: Google Custom Search API keys are not configured. Using simulated fallback.")
        # Fallback to simulated responses if API keys are not set
        if "Koppal Lok Sabha" in query or "MP Koppal" in query or "Member of Parliament Koppal" in query:
            return """
            (Simulated) Latest news (June 2024): K. Rajashekar Basavaraj Hitnal (Indian National Congress) won the Koppal Lok Sabha constituency in the 2024 General Elections. He defeated Basavaraj Kyavator (Bharatiya Janata Party).
            Source: Election Commission of India results, leading news outlets.
            """
        elif "Gangavati Constituency MLA" in query or "MLA of Gangavati" in query:
            return """
            (Simulated) Latest news (May 2023 Karnataka Assembly Elections): G. Janardhana Reddy (Kalyana Rajya Pragathi Paksha) is the current MLA for Gangavati Assembly Constituency.
            Source: Election Commission of India results, state assembly websites.
            """
        elif "today's weather" in query or "weather today" in query or "weather in Bengaluru" in query:
            return "(Simulated) The weather today in Bengaluru is partly cloudy with a high of 30°C and a low of 22°C. There's a 20% chance of rain."
        elif "prime minister of india" in query.lower():
            return "(Simulated) The current Prime Minister of India is Narendra Modi, who assumed office on May 26, 2014."
        elif "current president of india" in query.lower() or "who is the president of india" in query.lower():
            return "(Simulated) The current President of India is Droupadi Murmu. She assumed office on 25 July 2022."
        else:
            return f"(Simulated) No specific real-time data found for '{query}'. This is a simulated search result because API keys are not configured."

    search_url = f"https://www.googleapis.com/customsearch/v1?key={GOOGLE_CSE_API_KEY}&cx={GOOGLE_CSE_ID}&q={query}"

    st.info(f"🔍 Performing real web search for: '{query}' using Google Custom Search API...")
    try:
        response = requests.get(search_url)
        response.raise_for_status() # Raise an HTTPError for bad responses (4xx or 5xx)
        search_results = response.json()

        if search_results and 'items' in search_results:
            snippets = []
            for item in search_results['items'][:3]: # Get top 3 snippets for brevity
                title = item.get('title', 'No Title')
                snippet = item.get('snippet', 'No Snippet')
                link = item.get('link', '#')
                snippets.append(f"Title: {title}\nSnippet: {snippet}\nLink: {link}")
            return "\n---\n".join(snippets)
        else:
            return f"No relevant search results found for '{query}' from real API."
    except requests.exceptions.RequestException as e:
        st.error(f"❌ Error during real search API call: {e}")
        return f"Error connecting to real search service: {e}. Check your API keys and network."
    except Exception as e:
        st.error(f"❌ An unexpected error occurred during real search: {e}")
        return f"An unexpected error occurred during search: {e}"

# --- Custom Document Content Loading (Now supports .txt, .pdf, .docx) ---
def read_document_content_from_file(uploaded_file):
    """
    Reads the content of an uploaded file (text, PDF, or DOCX) and returns it as a string.
    """
    file_extension = os.path.splitext(uploaded_file.name)[1].lower()
    content = ""

    try:
        if file_extension == '.txt':
            content = uploaded_file.read().decode('utf-8')
        elif file_extension == '.pdf':
            # Save to a temporary file to allow pypdf to read it
            with open("temp_uploaded.pdf", "wb") as f:
                f.write(uploaded_file.getbuffer())
            try:
                with open("temp_uploaded.pdf", 'rb') as f:
                    reader = pypdf.PdfReader(f)
                    for page_num in range(len(reader.pages)):
                        page_text = reader.pages[page_num].extract_text()
                        if page_text:
                            content += page_text + "\n"
                if not content.strip():
                    st.warning("❗ Warning: PDF seemed to contain no extractable text. This often happens with scanned PDFs (images of text). OCR is not supported in this version.")
                    return None
            except pypdf.errors.PdfReadError:
                st.error(f"❌ Error reading PDF: It might be corrupted or encrypted.")
                return None
            finally:
                os.remove("temp_uploaded.pdf") # Clean up temp file
        elif file_extension == '.docx':
            # Save to a temporary file to allow python-docx to read it
            with open("temp_uploaded.docx", "wb") as f:
                f.write(uploaded_file.getbuffer())
            try:
                doc = Document("temp_uploaded.docx")
                for paragraph in doc.paragraphs:
                    content += paragraph.text + "\n"
            except Exception as docx_e:
                st.error(f"❌ Error reading DOCX: {docx_e}. Ensure it's a valid .docx file.")
                return None
            finally:
                os.remove("temp_uploaded.docx") # Clean up temp file
        else:
            st.error(f"❌ Error: Unsupported file type '{file_extension}'. Only .txt, .pdf, and .docx are supported.")
            return None

        st.success(f"📄 Successfully loaded content from: {uploaded_file.name}")
        return content
    except Exception as e:
        st.error(f"❌ An unexpected error occurred while reading file: {e}")
        return None

# --- Local Flan-T5 Answer Generation (Contextual QA) ---
def get_local_answer(question_text, context_text=None):
    """
    Generates an answer to a question using a local Flan-T5 model,
    optionally using provided context.
    """
    if local_qa_generator is None:
        return "Local QA model is not loaded. Cannot generate answer."

    st.info("🤔 Generating answer with local model (please wait)...")
    try:
        if context_text:
            prompt = f"Provide a comprehensive answer to the question based ONLY on the provided context. Context: {context_text}\n\nQuestion: {question_text}"
        else:
            # This path should ideally not be taken if routing is correct, but as a fallback.
            prompt = f"Answer this question clearly and accurately: {question_text}"

        answer_output = local_qa_generator(prompt, max_new_tokens=128) # Limit output to 128 tokens
        answer = answer_output[0]['generated_text'].strip()
        st.info(f"DEBUG: Local model generated: '{answer}' (Length: {len(answer)})") # Debug print
        
        # Add explicit check for empty or very short/non-sensical answers from local model
        # This is the key to force fallback if Flan-T5 fails to provide a good answer
        if not answer or len(answer) < 5 or "i cannot answer" in answer.lower() or "not in the response" in answer.lower() or "sorry" in answer.lower() or "could not find the answer" in answer.lower():
            st.warning("❗ Local model produced a short/unhelpful answer. Will try Gemini.")
            return "LOCAL_MODEL_FALLBACK_TRIGGERED" # Special signal to trigger Gemini fallback
        
        return answer
    except Exception as e:
        st.error(f"❌ Error during local model inference: {e}")
        return "LOCAL_MODEL_FALLBACK_TRIGGERED" # Trigger fallback on error as well

# --- Emotion Detection Function ---
def detect_emotion(text: str):
    """
    Detects the emotion in the given text using the pre-loaded emotion classifier.
    Returns the detected label and its score.
    """
    if emotion_classifier is None:
        return {"label": "unknown", "score": 0.0, "message": "Emotion classifier not loaded."}
    try:
        result = emotion_classifier(text)[0]
        return {
            "label": result['label'],
            "score": round(result['score'], 2),
            "message": f"Detected Emotion: {result['label']} (Score: {round(result['score'], 2)})"
        }
    except Exception as e:
        st.error(f"❌ Error detecting emotion: {e}")
        return {"label": "error", "score": 0.0, "message": f"Error detecting emotion: {e}"}

# --- GEMINI API Answer Generation (General QA, Tool Use, Conversational Memory) ---
def get_gemini_answer(question_text: str):
    """
    Fetches an answer from the Gemini 2.0 Flash API, incorporating tool use
    for real-time information and conversational memory.
    """
    # Use st.session_state for chat_history
    # User input is appended at the top of the user_input block, so not here.

    tools = [
        {
            "functionDeclarations": [
                {
                    "name": "search_tool",
                    "description": "Performs a real-time web search for current events, facts, or dynamic information that is not available in the model's training data or provided context.",
                    "parameters": {
                        "type": "OBJECT",
                        "properties": {
                            "query": {
                                "type": "STRING",
                                "description": "The search query to execute to get up-to-date information."
                            }
                        },
                        "required": ["query"]
                    }
                }
            ]
        }
    ]

    payload = {
        "contents": st.session_state.chat_history,
        "tools": tools,
        "generationConfig": {
            "temperature": 0.9,
            "maxOutputTokens": 800
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("🤔 Getting answer from Gemini..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()

            if result.get("candidates") and result["candidates"][0].get("content") and \
               result["candidates"][0]["content"].get("parts") and \
               result["candidates"][0]["content"]["parts"][0].get("functionCall"):
                
                function_call = result["candidates"][0]["content"]["parts"][0]["functionCall"]
                function_name = function_call["name"]
                function_args = function_call["args"]

                st.info(f"🤖 Gemini decided to call a tool: {function_name} with args {function_args}")

                st.session_state.chat_history.append({"role": "model", "parts": [{"functionCall": function_call}]})
                
                if function_name == "search_tool":
                    tool_response_content = search_tool(function_args.get("query")) # Call sync search_tool
                    
                    st.session_state.chat_history.append({
                        "role": "tool",
                        "parts": [{"functionResponse": {"name": function_name, "response": {"text": tool_response_content}}}]
                    })

                    payload_with_tool_response = {
                        "contents": st.session_state.chat_history,
                        "tools": tools,
                        "generationConfig": payload["generationConfig"]
                    }
                    response_after_tool = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload_with_tool_response))
                    response_after_tool.raise_for_status()
                    result_after_tool = response_after_tool.json()

                    if result_after_tool.get("candidates") and result_after_tool["candidates"][0].get("content") and \
                       result_after_tool["candidates"][0]["content"].get("parts"):
                        final_answer = result_after_tool["candidates"][0]["content"]["parts"][0]["text"]
                        st.session_state.chat_history.append({"role": "model", "parts": [{"text": final_answer}]})
                        st.info(f"DEBUG: Gemini (after tool) generated: {final_answer}") # Debug print
                        return final_answer
                    else:
                        st.warning("❗ Warning: Gemini API (after tool) response structure unexpected or content missing.")
                        return "Sorry, I processed the search but couldn't get a clear answer from Gemini."
                else:
                    return "Sorry, Gemini tried to call an unknown tool."

            if result.get("candidates") and len(result["candidates"]) > 0 and \
               result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                direct_answer = result["candidates"][0]["content"].get("parts")[0].get("text", "No text found in response.")
                st.session_state.chat_history.append({"role": "model", "parts": [{"text": direct_answer}]})
                st.info(f"DEBUG: Gemini (direct) generated: {direct_answer}") # Debug print
                return direct_answer
            else:
                st.warning("❗ Warning: Gemini API response structure unexpected or content missing.")
                return "Sorry, I couldn't get a direct answer from Gemini."

        except requests.exceptions.RequestException as e:
            st.error(f"❌ Error calling Gemini API: {e}. Check your internet connection or API key.")
            return "Sorry, I'm having trouble connecting to the answer service."
        except Exception as e:
            st.error(f"❌ An unexpected error occurred with Gemini: {e}")
            return "An unexpected error occurred while getting the answer from Gemini."

# --- New Feature: Quiz Generation ---
def generate_quiz(topic: str, context: str):
    """Generates quiz questions based on a topic and provided context using Gemini."""
    st.info(f"📝 Generating a quiz on '{topic}' from the provided notes...")
    quiz_prompt = f"""
    You are an educational assistant. Generate 5-7 multiple-choice quiz questions (with 4 options each) and their correct answers based on the following text about {topic}.
    Ensure the questions and options are clear and directly related to the text. Provide the correct answer letter.

    Text:
    {context}

    Format the output as follows:

    Question 1: [Question text]
    A) [Option A]
    B) [Option B]
    C) [Option C]
    D) [Option D]
    Correct Answer: [A/B/C/D]

    Question 2: ...
    """
    quiz_history_temp = [{"role": "user", "parts": [{"text": quiz_prompt}]}]
    
    payload = {
        "contents": quiz_history_temp,
        "generationConfig": {
            "temperature": 0.7,
            "maxOutputTokens": 800
        }
    }
    
    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Generating quiz..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                quiz_text = result["candidates"][0]["content"].get("parts")[0].get("text", "No quiz generated.")
                st.session_state.last_generated_quiz = quiz_text
                return quiz_text
            else:
                st.session_state.last_generated_quiz = None
                return "Could not generate quiz. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            st.session_state.last_generated_quiz = None
            return f"Error generating quiz: {e}"
        except Exception as e:
            st.session_state.last_generated_quiz = None
            return f"An unexpected error occurred during quiz generation: {e}"

# --- New Feature: Identify Important Questions from Quiz ---
def identify_important_questions_from_quiz(quiz_text: str):
    """Identifies and extracts important questions from a given quiz text using Gemini."""
    st.info("🧠 Identifying important questions from the quiz...")
    importance_prompt = f"""
    From the following quiz questions, identify the 2 to 3 most important or foundational questions.
    Present them as a numbered list.
    
    Quiz:
    {quiz_text}
    """
    importance_history_temp = [{"role": "user", "parts": [{"text": importance_prompt}]}]

    payload = {
        "contents": importance_history_temp,
        "generationConfig": {
            "temperature": 0.5,
            "maxOutputTokens": 300
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Identifying important questions..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                important_questions = result["candidates"][0]["content"].get("parts")[0].get("text", "Could not identify important questions.")
                return important_questions
            else:
                return "Could not identify important questions. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            return f"Error identifying important questions: {e}"
        except Exception as e:
            return f"An unexpected error occurred while identifying important questions: {e}"

# NEW: Identify Important Points from Summary
def identify_important_points_from_summary(summary_text: str):
    """Identifies and extracts important points from a given summary text using Gemini."""
    st.info("🧠 Identifying important points from the summary...")
    importance_prompt = f"""
    From the following summary, extract the 3 to 5 most crucial key points.
    Present them as a numbered list.

    Summary:
    {summary_text}
    """
    importance_history_temp = [{"role": "user", "parts": [{"text": importance_prompt}]}]

    payload = {
        "contents": importance_history_temp,
        "generationConfig": {
            "temperature": 0.3,
            "maxOutputTokens": 200
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Identifying important points..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                important_points = result["candidates"][0]["content"].get("parts")[0].get("text", "Could not identify important points.")
                return important_points
            else:
                return "Could not identify important points. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            return f"Error identifying important points: {e}"
        except Exception as e:
            return f"An unexpected error occurred while identifying important points: {e}"

# NEW: Generate Important Questions from Notes
def generate_important_questions_from_notes(notes_content: str):
    """Generates important questions directly from the full notes content using Gemini."""
    st.info("🧠 Generating important questions directly from the notes...")
    questions_prompt = f"""
    You are an educational assistant. Based on the following study notes, generate 3 to 5 comprehensive and important open-ended questions that a student should be able to answer after studying these notes. These questions should cover the most crucial concepts.
    Present them as a numbered list.

    Notes:
    {notes_content}
    """
    questions_history_temp = [{"role": "user", "parts": [{"text": questions_prompt}]}]

    payload = {
        "contents": questions_history_temp,
        "generationConfig": {
            "temperature": 0.6,
            "maxOutputTokens": 400
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Generating important questions..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                important_questions = result["candidates"][0]["content"].get("parts")[0].get("text", "Could not generate important questions from notes.")
                return important_questions
            else:
                return "Could not generate important questions from notes. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            return f"Error generating important questions from notes: {e}"
        except Exception as e:
            return f"An unexpected error occurred while generating important questions from notes: {e}"

# NEW: Generate Practice Questions from Notes
def generate_practice_questions(notes_content: str):
    """Generates general practice questions directly from the full notes content using Gemini."""
    st.info("🧠 Generating practice questions directly from the notes...")
    questions_prompt = f"""
    You are an educational assistant. Based on the following study notes, generate 3 to 5 comprehensive practice questions that cover various aspects of the material. These should be good questions for a student to use for self-study.
    Present them as a numbered list.

    Notes:
    {notes_content}
    """
    questions_history_temp = [{"role": "user", "parts": [{"text": questions_prompt}]}]

    payload = {
        "contents": questions_history_temp,
        "generationConfig": {
            "temperature": 0.7,
            "maxOutputTokens": 400
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Generating practice questions..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                practice_questions = result["candidates"][0]["content"].get("parts")[0].get("text", "Could not generate practice questions from notes.")
                return practice_questions
            else:
                return "Could not generate practice questions from notes. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            return f"Error generating practice questions from notes: {e}"
        except Exception as e:
            return f"An unexpected error occurred while generating practice questions from notes: {e}"


# --- New Feature: Content Summarization ---
def summarize_content(text_to_summarize: str):
    """Summarizes provided text content using Gemini."""
    st.info("📚 Summarizing the notes...")
    summary_prompt = f"""
    Please provide a concise and informative summary of the following text using a numbered list (1., 2., 3., etc.).
    Each numbered point should highlight a main point or key takeaway.

    Text:
    {text_to_summarize}
    """
    summary_history_temp = [{"role": "user", "parts": [{"text": summary_prompt}]}]

    payload = {
        "contents": summary_history_temp,
        "generationConfig": {
            "temperature": 0.3,
            "maxOutputTokens": 500
        }
    }

    apiUrl = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={GEMINI_API_KEY}"

    with st.spinner("Generating summary..."):
        try:
            response = requests.post(apiUrl, headers={'Content-Type': 'application/json'}, data=json.dumps(payload))
            response.raise_for_status()
            result = response.json()
            if result.get("candidates") and result["candidates"][0].get("content") and result["candidates"][0]["content"].get("parts"):
                summary_text = result["candidates"][0]["content"].get("parts")[0].get("text", "No summary generated.")
                st.session_state.last_generated_summary = summary_text
                return summary_text
            else:
                st.session_state.last_generated_summary = None
                return "Could not generate summary. Gemini API response unexpected."
        except requests.exceptions.RequestException as e:
            st.session_state.last_generated_summary = None
            return f"Error summarizing content: {e}"
        except Exception as e:
            st.session_state.last_generated_summary = None
            return f"An unexpected error occurred during summarization: {e}"

# --- Document Chunking and Retrieval (for Local Model) ---
def chunk_text(text, chunk_size=400, overlap=50): # Reduced chunk_size slightly
    """Splits text into chunks of specified size with overlap."""
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size - overlap):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
    return chunks

def find_relevant_chunks(question, chunks, top_n=2):
    """
    Finds chunks most relevant to the question using simple keyword matching.
    Returns a concatenated string of the most relevant chunks.
    """
    question_keywords = set(word.lower() for word in question.split() if len(word) > 2)
    
    scored_chunks = []
    for i, chunk in enumerate(chunks):
        chunk_keywords = set(word.lower() for word in chunk.split() if len(word) > 2)
        score = len(question_keywords.intersection(chunk_keywords))
        if score > 0:
            scored_chunks.append((score, i, chunk))
    
    scored_chunks.sort(key=lambda x: x[0], reverse=True)
    
    relevant_chunks = [chunk for score, i, chunk in scored_chunks[:top_n]]

    if not relevant_chunks:
        # If no chunks match keywords, return an empty string to signal no relevant content
        return "" 
    
    return "\n---\n".join(relevant_chunks) # Join with a separator

# --- Streamlit UI Layout ---
st.set_page_config(page_title="Smart Classroom Assistant", layout="centered")

st.title("📚 Smart Classroom Assistant")
st.markdown("Your AI-powered study companion for notes, quizzes, and general knowledge.")

# File Uploader for Notes
st.sidebar.header("Upload Your Notes")
uploaded_file = st.sidebar.file_uploader("Choose a TXT, PDF, or DOCX file", type=["txt", "pdf", "docx"])

if uploaded_file is not None:
    # Check if a new file is uploaded or if it's a rerun with the same file
    if st.session_state.loaded_file_name != uploaded_file.name:
        st.session_state.science_notes_content = read_document_content_from_file(uploaded_file)
        if st.session_state.science_notes_content:
            st.session_state.loaded_file_name = uploaded_file.name
            st.session_state.science_notes_chunks = chunk_text(st.session_state.science_notes_content)
            st.session_state.chat_history = [] # Clear chat history on new document load
            st.sidebar.success(f"Notes loaded: {uploaded_file.name}")
        else:
            st.session_state.loaded_file_name = None
            st.session_state.science_notes_content = None
            st.session_state.science_notes_chunks = []
            st.sidebar.error("Failed to load notes.")
    else:
        st.sidebar.info(f"Notes already loaded: {st.session_state.loaded_file_name}")
else:
    # This block handles when the user removes the file by clicking 'x'
    if st.session_state.loaded_file_name is not None: # Only reset if a file was previously loaded
        st.session_state.loaded_file_name = None
        st.session_state.science_notes_content = None
        st.session_state.science_notes_chunks = []
        st.session_state.chat_history = [] # Clear chat history on document removal
        st.sidebar.info("Notes removed. Assistant will now answer general questions.")


if st.session_state.loaded_file_name:
    st.sidebar.markdown(f"**Current Notes:** {st.session_state.loaded_file_name}")
else:
    st.sidebar.info("No notes loaded. Assistant will answer general questions.")

# --- Command Buttons ---
st.sidebar.header("Quick Actions")
col1, col2 = st.sidebar.columns(2)

with col1:
    if st.button("Summarize Notes", key="summarize_btn"):
        if st.session_state.science_notes_content:
            summary_output = summarize_content(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Summarize notes"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": summary_output}]})
        else:
            st.error("Please upload notes first to summarize.")

    if st.button("Generate Quiz", key="quiz_btn"):
        if st.session_state.science_notes_content:
            quiz_topic = "Your Notes"
            if st.session_state.loaded_file_name:
                base_name = os.path.splitext(st.session_state.loaded_file_name)[0]
                quiz_topic = base_name.replace('_', ' ').replace('-', ' ').title()
            quiz_output = generate_quiz(quiz_topic, st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Generate quiz"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": quiz_output}]})
        else:
            st.error("Please upload notes first to generate a quiz.")

    if st.button("Important Qs (Notes)", key="imp_q_notes_btn"):
        if st.session_state.science_notes_content:
            imp_q_notes_output = generate_important_questions_from_notes(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Important questions from notes"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_q_notes_output}]})
        else:
            st.error("Please upload notes first to get important questions.")

with col2:
    if st.button("Practice Questions", key="practice_q_btn"):
        if st.session_state.science_notes_content:
            practice_q_output = generate_practice_questions(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Generate practice questions"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": practice_q_output}]})
        else:
            st.error("Please upload notes first to generate practice questions.")

    if st.button("Important Qs (Quiz)", key="imp_q_quiz_btn"):
        if st.session_state.last_generated_quiz:
            imp_q_quiz_output = identify_important_questions_from_quiz(st.session_state.last_generated_quiz)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Important questions from last quiz"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_q_quiz_output}]})
        else:
            st.warning("No quiz has been generated yet. Please generate a quiz first.")

    if st.button("Important Pts (Summary)", key="imp_p_summary_btn"):
        if st.session_state.last_generated_summary:
            imp_p_summary_output = identify_important_points_from_summary(st.session_state.last_generated_summary)
            st.session_state.chat_history.append({"role": "user", "parts": [{"text": "Important points from last summary"}]})
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_p_summary_output}]})
        else:
            st.warning("No summary has been generated yet. Please summarize notes first.")

if st.sidebar.button("Clear Chat History", key="clear_chat_btn"):
    st.session_state.chat_history = []
    st.success("Chat history cleared.")

# --- Chat Display Area ---
st.subheader("Conversation")
chat_container = st.container(height=400, border=True)

for message in st.session_state.chat_history:
    if message["role"] == "user":
        chat_container.chat_message("user").write(message["parts"][0]["text"])
    elif message["role"] == "model":
        # Check if the model's part is a text response or a function call
        part = message["parts"][0]
        if "text" in part:
            chat_container.chat_message("assistant").write(part["text"])
        elif "functionCall" in part:
            function_call_info = part["functionCall"]
            chat_container.chat_message("assistant").info(f"Assistant called a tool: {function_call_info['name']}({function_call_info['args']})")
        else:
            # Fallback for unexpected model content structure
            chat_container.chat_message("assistant").write("Assistant: [Unexpected message format]")
    elif message["role"] == "tool":
        # For tool responses, display them clearly
        tool_response_content = message["parts"][0]["functionResponse"]["response"]["text"]
        chat_container.chat_message("assistant").info(f"Tool Response: {tool_response_content}")

# --- User Input ---
user_input = st.chat_input("Ask a question or type a command...")

if user_input:
    # Always append user input first
    st.session_state.chat_history.append({"role": "user", "parts": [{"text": user_input}]})
    
    # Detect emotion
    emotion_result = detect_emotion(user_input)
    st.info(f"Detected Emotion: {emotion_result['label']} (Score: {emotion_result['score']})")

    # Command Handling (simplified for Streamlit)
    cleaned_input_text = user_input.lower().strip()

    # Check for specific commands first
    if cleaned_input_text == "clear chat":
        st.session_state.chat_history = [] # This will trigger a rerun and clear display
        st.success("Chat history cleared.")
        st.rerun() 
    elif cleaned_input_text == "show history":
        # History is already displayed in the chat_container, so just acknowledge
        st.session_state.chat_history.append({"role": "model", "parts": [{"text": "The conversation history is displayed above."}]})
        st.rerun()
    elif cleaned_input_text == "exit":
        st.session_state.chat_history.append({"role": "model", "parts": [{"text": "Assistant session ended. You can close this tab."}]})
        st.stop() 
    elif cleaned_input_text == "summarize notes":
        if st.session_state.science_notes_content:
            summary_output = summarize_content(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": summary_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "Please upload notes first to summarize."}]})
        st.rerun()
    elif cleaned_input_text == "generate quiz": 
        if st.session_state.science_notes_content:
            quiz_topic = "Your Notes"
            if st.session_state.loaded_file_name:
                base_name = os.path.splitext(st.session_state.loaded_file_name)[0]
                quiz_topic = base_name.replace('_', ' ').replace('-', ' ').title()
            quiz_output = generate_quiz(quiz_topic, st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": quiz_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "Please upload notes first to generate a quiz."}]})
        st.rerun()
    elif cleaned_input_text == "important questions from quiz": 
        if st.session_state.last_generated_quiz:
            imp_q_output = identify_important_questions_from_quiz(st.session_state.last_generated_quiz)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_q_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "No quiz has been generated yet. Please use 'Generate Quiz' first."}]})
        st.rerun()
    elif cleaned_input_text == "important points from summary": 
        if st.session_state.last_generated_summary:
            imp_p_output = identify_important_points_from_summary(st.session_state.last_generated_summary)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_p_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "No summary has been generated yet. Please use 'Summarize Notes' first."}]})
        st.rerun()
    elif cleaned_input_text == "important questions from notes": 
        if st.session_state.science_notes_content:
            imp_q_notes_output = generate_important_questions_from_notes(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": imp_q_notes_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "Please upload notes first to get important questions."}]})
        st.rerun()
    elif cleaned_input_text == "practice questions":
        if st.session_state.science_notes_content:
            practice_q_output = generate_practice_questions(st.session_state.science_notes_content)
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": practice_q_output}]})
        else:
            st.session_state.chat_history.append({"role": "model", "parts": [{"text": "Please upload notes first to generate practice questions."}]})
        st.rerun()
    # Add more command handling here if needed

    # If not a command, then it's a Q&A
    else:
        assistant_response = "" 
        if st.session_state.science_notes_content:
            st.info("🧠 Routing to local model for contextual Q&A with loaded notes (using chunks)...")
            relevant_context = find_relevant_chunks(user_input, st.session_state.science_notes_chunks)
            
            # Only use local model if relevant context is found AND local_qa_generator is loaded
            if relevant_context and local_qa_generator:
                local_answer = get_local_answer(user_input, context_text=relevant_context)
                
                # If local model's answer is a "sorry" type response or our special fallback trigger
                if local_answer == "LOCAL_MODEL_FALLBACK_TRIGGERED":
                    st.info("🧠 Local model couldn't answer from context or gave a poor response. Routing to Gemini for general knowledge/search...")
                    assistant_response = get_gemini_answer(user_input) # Gemini appends to chat_history internally
                else:
                    # Local model provided a good answer, append it directly
                    assistant_response = local_answer
                    # The local_answer is already appended in get_local_answer, but if it's not and it's a valid answer, append it here.
                    # To avoid double appending, we need to be careful.
                    # Let's ensure get_local_answer *does not* append to chat_history.
                    # And only the main loop appends.
                    # Re-checking get_local_answer: it returns a string, doesn't append. Good.
                    st.session_state.chat_history.append({"role": "model", "parts": [{"text": assistant_response}]})
            else:
                st.info("🧠 No relevant chunks found in the document or local model not loaded. Routing to Gemini for general knowledge/search...")
                assistant_response = get_gemini_answer(user_input) # Gemini appends to chat_history internally
        else:
            # If no document is loaded, always use Gemini for general knowledge
            st.info("🌐 Routing to Gemini for general Q&A / web search (no notes loaded)...")
            assistant_response = get_gemini_answer(user_input) # Gemini appends to chat_history internally
        
        st.rerun() 